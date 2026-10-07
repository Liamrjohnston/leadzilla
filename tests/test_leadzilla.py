"""Offline tests: no network, no Jev calls. Jev is replaced by a fake that returns fixed answers."""
import json
import tempfile
import unittest
from pathlib import Path

from leadzilla import qualify as Q
from leadzilla.extract import contact_links, external_sites, extract, phone_from_any
from leadzilla.jev import Answer, Meter
from leadzilla.pipeline import _brand, _names_from_email
from leadzilla.store import Store

PAGE = """<html><head><title>Acme Realty | Vancouver Brokerage</title>
<meta name="description" content="Boutique brokerage">
<script type="application/ld+json">{"@graph":[{"@type":"RealEstateAgent","name":"Acme Realty Ltd",
"telephone":"604-555-0199","email":"info@acmerealty.ca",
"address":{"streetAddress":"1 Main St","addressLocality":"Vancouver","addressRegion":"BC"}}]}</script>
<script>var x="tracker@sentry.io";</script></head>
<body><a href="mailto:Jane.Doe@AcmeRealty.ca?subject=hi">Email Jane</a><a href="tel:+1 (604) 555-0100">call</a>
<a href="/contact-us">Contact</a><a href="/listings">Listings</a><a href="https://www.instagram.com/acme">IG</a>
<a href="https://othersite.com/x">Partner</a><div>Office 604-629-6100</div><div>admin@acmerealty.ca</div>
<p>sales [at] acmerealty [dot] ca · logo@2x.png · 900-112-2970</p></body></html>"""


class Extraction(unittest.TestCase):
    def setUp(self):
        self.p = extract(PAGE, "https://www.acmerealty.ca/")

    def test_emails_from_every_source_and_junk_dropped(self):
        self.assertEqual(set(self.p["emails"]), {"jane.doe@acmerealty.ca", "info@acmerealty.ca",
                                                 "admin@acmerealty.ca", "sales@acmerealty.ca"})

    def test_script_contents_ignored(self):
        self.assertNotIn("tracker@sentry.io", self.p["emails"])

    def test_phones_normalised_and_fake_rejected(self):
        self.assertIn("(604) 555-0100", self.p["phones"])
        self.assertIn("(604) 629-6100", self.p["phones"])
        self.assertNotIn("(900) 112-2970", self.p["phones"])

    def test_jsonld_graph(self):
        self.assertEqual(self.p["jsonld"][0]["name"], "Acme Realty Ltd")
        self.assertIn("Vancouver", self.p["jsonld"][0]["address"])

    def test_links(self):
        self.assertEqual([l["url"] for l in contact_links(self.p)], ["https://www.acmerealty.ca/contact-us"])
        self.assertEqual([l["domain"] for l in external_sites(self.p)], ["othersite.com"])
        self.assertEqual(self.p["socials"]["instagram"], "https://www.instagram.com/acme")

    def test_bad_html_does_not_crash(self):
        self.assertEqual(extract("", "https://x.ca")["emails"], [])

    def test_phone_from_any(self):
        self.assertEqual(phone_from_any("+1-778-374-3100"), "(778) 374-3100")
        self.assertIsNone(phone_from_any("12345"))


class FakeJev:
    def __init__(self, answers):
        self.answers, self.meter, self.calls = answers, Meter(), []

    def ask(self, state, questions):
        self.calls.append((state, questions))
        return {k: self.answers(k, q) for k, q in questions.items()}


def lead(domain, emails):
    return {"domain": domain, "company": domain, "page_title": "", "description": "", "jsonld": [],
            "address": "", "excerpt": "text", "emails": emails}


class Qualification(unittest.TestCase):
    def fake(self, k, q):
        if k.endswith("_biz"): return Answer("noul", 0.9, None, {})
        if k.endswith("_area"): return Answer("noul", 0.8, None, {})
        if k.endswith("_fit"): return Answer("score", 3.2, 0.7, {})
        if k.endswith("_size"): return Answer("choice", "small", 0.9, {})
        if k.endswith("_email"): return Answer("choice", "e1", 0.8, {})
        raise AssertionError(k)

    def test_batching_and_email_pick(self):
        leads = [lead(f"d{i}.ca", ["a@d.ca", "b@d.ca"]) for i in range(7)]
        jev = FakeJev(self.fake)
        Q.judge_leads(jev, leads, "brokerages", "Vancouver", "offer", batch=5)
        self.assertEqual(len(jev.calls), 2)
        self.assertEqual(leads[0]["email"], "b@d.ca")
        self.assertEqual(Q.keep(leads[0])[0], "qualified")

    def test_single_email_needs_no_question(self):
        L = [lead("x.ca", ["only@x.ca"])]
        jev = FakeJev(self.fake)
        Q.judge_leads(jev, L, "n", "p", "o")
        self.assertNotIn("x0_email", jev.calls[0][1])
        self.assertEqual(L[0]["email"], "only@x.ca")

    def test_thresholds(self):
        base = {"is_business": 0.9, "in_area": 0.9, "fit": 3.0, "confidence": 0.8, "size": "small"}
        self.assertEqual(Q.keep({**base, "is_business": 0.3})[0], "rejected")
        self.assertEqual(Q.keep({**base, "is_business": 0.5})[0], "review")
        self.assertEqual(Q.keep({**base, "in_area": 0.2})[0], "rejected")
        self.assertEqual(Q.keep({**base, "fit": 1.5})[0], "rejected")

    def test_map_leads_skip_the_area_question(self):
        L = [{**lead("m.ca", []), "on_map": True}]
        jev = FakeJev(self.fake)
        Q.judge_leads(jev, L, "n", "p", "o")
        self.assertNotIn("x0_area", jev.calls[0][1])
        self.assertEqual(L[0]["in_area"], 1.0)

    def test_reason_is_built_by_code(self):
        r = Q.reason({"fit": 3.04, "confidence": 0.91, "size": "mid", "in_area": 0.97, "email": "i@x.ca",
                      "email_role": True})
        self.assertEqual(r, "Fit 3.0/4 (confidence 91%); mid-size firm; in area 97%; general inbox")


class Memory(unittest.TestCase):
    def test_resume_and_dedupe(self):
        with tempfile.TemporaryDirectory() as d:
            s = Store(Path(d) / "lz.db")
            cid = s.campaign("Brokerages", "Vancouver", "Offer")
            self.assertEqual(cid, s.campaign("brokerages ", "vancouver", "offer"))
            s.mark(cid, "a.ca", "rejected", "x")
            self.assertTrue(s.seen(cid, "a.ca"))
            s.add_lead(cid, {"domain": "b.ca", "fit": 3.0, "socials": {}, "data": {}})
            s.add_lead(cid, {"domain": "b.ca", "fit": 3.5, "socials": {}, "data": {}})
            self.assertEqual(s.lead_count(cid), 1)


class Export(unittest.TestCase):
    LEADS = [
        {"company": "Engel & Völkers", "website": "https://a.ca/", "email": "info@a.ca", "email_role": 1,
         "email_status": "domain_accepts_mail", "phone": "(604) 555-0100", "address": "", "fit": 3.1,
         "confidence": 0.8, "size": "mid", "reason": "r", "source": "map",
         "socials": json.dumps({"linkedin": "https://www.linkedin.com/company/a"}),
         "data": json.dumps({"all_emails": ["info@a.ca", "b@a.ca"], "status": "qualified"})},
        {"company": "No Email Co", "website": "https://b.ca/", "email": None, "email_role": 0,
         "email_status": None, "phone": None, "address": "", "fit": 2.5, "confidence": 0.6, "size": "small",
         "reason": "r", "source": "map", "socials": "{}", "data": json.dumps({"status": "review"})},
    ]

    def test_csv_excel_safe_and_complete(self):
        from leadzilla import export as E
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "l.csv"
            self.assertEqual(E.write(self.LEADS, str(p), "csv"), 2)
            raw = p.read_bytes()
            self.assertTrue(raw.startswith(b"\xef\xbb\xbf"))
            text = raw.decode("utf-8-sig")
            self.assertIn("Engel & Völkers", text)
            self.assertIn("b@a.ca", text)

    def test_email_tool_presets_drop_rows_without_email(self):
        from leadzilla import export as E
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "i.csv"
            self.assertEqual(E.write(self.LEADS, str(p), "instantly"), 1)
            self.assertTrue(p.read_text(encoding="utf-8-sig").startswith("email,first_name"))

    def test_xlsx(self):
        from leadzilla import export as E
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            p = Path(d) / "l.xlsx"
            E.write(self.LEADS, str(p), "xlsx")
            ws = load_workbook(p).active
            self.assertEqual(ws["A1"].value, "Status")
            self.assertEqual(ws.freeze_panes, "B2")
            self.assertTrue(ws["C2"].hyperlink.target.startswith("https://"))


class Helpers(unittest.TestCase):
    def test_brand(self):
        self.assertEqual(_brand("https://suttoncentrerealty.com"), _brand("https://www.suttoncentrerealty.ca/"))
        self.assertNotEqual(_brand("https://amyandally.com"), _brand("https://www.hugedomains.com/x"))
        self.assertEqual(_brand("https://shop.example.co.uk"), "example")

    def test_names(self):
        self.assertEqual(_names_from_email("jane.doe@x.ca"), ("Jane", "Doe"))
        self.assertEqual(_names_from_email("info@x.ca"), ("", ""))

    def test_meter_price(self):
        m = Meter(); m.add(1_000_000, 1.0)
        self.assertAlmostEqual(m.cost_usd, 0.042)


if __name__ == "__main__":
    unittest.main()
